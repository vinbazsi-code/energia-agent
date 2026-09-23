"""Közös segédfüggvények a MAVIR kiegyenlítő energia egységár gyűjtőkhöz
(napi "előzetes" és a historikus "végleges" visszatöltő is ezt használja).
"""

import io
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
from http_utils import get_with_retry  # noqa: E402

BASE = "https://mavir.hu"
USER_AGENT = "energia-agent-collector/0.1 (kontakt: vinbazsi@gmail.com)"
BUDAPEST_TZ = ZoneInfo("Europe/Budapest")

ELOZETES_ROOT = f"{BASE}/web/riportok/elozetes-kiegyenlito-energia-egysegarak"
VEGLEGES_ROOT = f"{BASE}/web/riportok/kiegyenlito-energia-egysegarak"


def http_get(url: str, log=None):
    return get_with_retry(url, headers={"User-Agent": USER_AGENT}, timeout=30)


def find_month_folder_url(root_url: str, month_key: str, log=None) -> str:
    """A gyökér dokumentumtár-listából kikeresi a `month_key` (pl. '202608') hónap mappájának 'view' URL-jét."""
    resp = http_get(root_url)
    pattern = re.compile(
        r'href="(' + re.escape(BASE) + r'/web/riportok/[^"]*?/view/(\d+)[^"]*)"[^>]*>\s*'
        + re.escape(month_key)
    )
    match = pattern.search(resp.text)
    if not match:
        raise RuntimeError(f"Nem található a(z) {month_key} hónap mappája a {root_url} oldalon")
    folder_url = match.group(1).replace("&amp;", "&")
    if log:
        log.info("Hónap mappa (%s) megtalálva: %s", month_key, folder_url)
    return folder_url


def find_latest_file_view_url(folder_url: str, filename_prefix: str, log=None) -> tuple[str, str]:
    """A hónap-mappa listából kikeresi a legfrissebb (első találat = legújabb) fájl view_file URL-jét."""
    resp = http_get(folder_url)
    pattern = re.compile(
        r'href="(' + re.escape(BASE) + r'/web/riportok/[^"]*?/view_file/(\d+)[^"]*)"[^>]*>\s*'
        + re.escape(filename_prefix) + r'([^<]*\.xlsx)'
    )
    match = pattern.search(resp.text)
    if not match:
        raise RuntimeError(f"Nem található '{filename_prefix}*.xlsx' fájl a {folder_url} mappában")
    view_file_url = match.group(1).replace("&amp;", "&")
    filename = filename_prefix + match.group(3)
    if log:
        log.info("Fájl megtalálva: %s", filename)
    return view_file_url, filename


def resolve_direct_download_url(view_file_url: str) -> str:
    """A fájl-előnézet oldal statikus HTML-jéből kinyeri a közvetlen /documents/... letöltési linket."""
    resp = http_get(view_file_url)
    match = re.search(r'https://mavir\.hu/documents/[^"\']+download=true', resp.text)
    if not match:
        raise RuntimeError(f"Nem található közvetlen letöltési link itt: {view_file_url}")
    return match.group(0).replace("&amp;", "&")


def download_xlsx(url: str) -> bytes:
    resp = http_get(url)
    content_type = resp.headers.get("Content-Type", "")
    if "spreadsheet" not in content_type:
        raise RuntimeError(f"Váratlan Content-Type ({content_type}) ettől: {url}")
    return resp.content


def parse_ke_file(xlsx_bytes: bytes, price_column_hint: str) -> dict[tuple[str, str], dict]:
    """Egy Poz_/Neg_ KE egységár Excel fájlt dolgoz fel.

    Visszaad egy {(dátum, intervallum): {rendszerirany_kwh, rendszerallapot_kwh, ar_huf_per_kwh}} dict-et.
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb.worksheets[0]

    header_row_idx = None
    header = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if row and row[0] == "Dátum":
            header_row_idx = i
            header = {str(name).strip(): idx for idx, name in enumerate(row) if name}
            break
    if header_row_idx is None:
        raise RuntimeError("Nem található a 'Dátum' fejléc a fájlban - megváltozott a formátum?")

    price_col = next((idx for name, idx in header.items() if price_column_hint in name), None)
    irany_col = header.get("Rendszer-irány (kWh)")
    allapot_col = header.get("Rendszerállapot (kWh)")
    if price_col is None or irany_col is None or allapot_col is None:
        raise RuntimeError(
            f"Hiányzó oszlop(ok) - price_col={price_col}, irany_col={irany_col}, allapot_col={allapot_col}"
        )

    results: dict[tuple[str, str], dict] = {}
    rows = list(ws.iter_rows(values_only=True))
    date_pattern = re.compile(r"^\d{4}-\d{2}-\d{2}$")
    for row in rows[header_row_idx + 1:]:
        if not row or not row[0] or not date_pattern.match(str(row[0])):
            continue
        datum = str(row[0])
        intervallum = str(row[1])
        results[(datum, intervallum)] = {
            "rendszerirany_kwh": row[irany_col],
            "rendszerallapot_kwh": row[allapot_col],
            "ar_huf_per_kwh": row[price_col],
        }
    return results


def local_interval_start_to_utc(datum: str, intervallum: str) -> datetime:
    """'2026-09-17' + '00:00 - 00:15' -> UTC datetime az intervallum kezdetéhez."""
    start_str = intervallum.split(" - ")[0].strip()
    local_dt = datetime.strptime(f"{datum} {start_str}", "%Y-%m-%d %H:%M").replace(tzinfo=BUDAPEST_TZ)
    return local_dt.astimezone(timezone.utc)


def fetch_and_parse_pair(folder_url: str, filename_prefix: str, log=None) -> tuple[dict, dict]:
    """Letölti és feldolgozza a Poz_/Neg_ fájlpárt egy adott hónap-mappából."""
    poz_view_url, poz_filename = find_latest_file_view_url(folder_url, f"Poz_{filename_prefix}", log)
    neg_view_url, neg_filename = find_latest_file_view_url(folder_url, f"Neg_{filename_prefix}", log)

    poz_download_url = resolve_direct_download_url(poz_view_url)
    neg_download_url = resolve_direct_download_url(neg_view_url)

    poz_bytes = download_xlsx(poz_download_url)
    neg_bytes = download_xlsx(neg_download_url)
    if log:
        log.info(
            "Letöltve: %s (%d KB), %s (%d KB)",
            poz_filename, len(poz_bytes) // 1024, neg_filename, len(neg_bytes) // 1024,
        )

    poz_data = parse_ke_file(poz_bytes, "Pozitív mérlegköri")
    neg_data = parse_ke_file(neg_bytes, "Negatív mérlegköri")
    return poz_data, neg_data


def build_measurement_rows(poz_data: dict, neg_data: dict, source_name: str, collected_at: datetime) -> list[dict]:
    """Poz_/Neg_ parse-eredményekből 'measurements' táblába illő sorokat épít."""
    all_keys = sorted(set(poz_data) | set(neg_data))
    collected_iso = collected_at.isoformat()
    rows = []
    for datum, intervallum in all_keys:
        utc_start = local_interval_start_to_utc(datum, intervallum)
        ts = utc_start.isoformat()
        poz = poz_data.get((datum, intervallum), {})
        neg = neg_data.get((datum, intervallum), {})
        rendszerirany = poz.get("rendszerirany_kwh", neg.get("rendszerirany_kwh"))
        rendszerallapot = poz.get("rendszerallapot_kwh", neg.get("rendszerallapot_kwh"))
        pozitiv_ar = poz.get("ar_huf_per_kwh")
        negativ_ar = neg.get("ar_huf_per_kwh")

        for metric, value, unit in (
            ("rendszerirany_kwh", rendszerirany, "kWh"),
            ("rendszerallapot_kwh", rendszerallapot, "kWh"),
            ("pozitiv_ar_huf_per_kwh", pozitiv_ar, "HUF/kWh"),
            ("negativ_ar_huf_per_kwh", negativ_ar, "HUF/kWh"),
        ):
            if value is None:
                continue
            rows.append(
                {
                    "timestamp_utc": ts,
                    "source": source_name,
                    "scope": "system",
                    "asset_id": "",
                    "metric": metric,
                    "value": value,
                    "unit": unit,
                    "collected_at": collected_iso,
                }
            )
    return rows
